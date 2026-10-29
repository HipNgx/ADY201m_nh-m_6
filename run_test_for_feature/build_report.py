from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

P=Path(__file__).resolve().parent
D=json.loads((P/'report_data.json').read_text(encoding='utf-8'))
S={(x['Experiment'],x['Model']):x for x in D['summary']}
protocol=D['protocol']
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.65);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Calibri')
for st in doc.styles:
 for border in list(st._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
normal=doc.styles['Normal'];normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(25);doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 1'].font.size=Pt(17);doc.styles['Heading 1'].paragraph_format.space_before=Pt(0);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].paragraph_format.space_before=Pt(9);doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.italic=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
rr=footer.add_run('Nhóm 6  |  Thí nghiệm bổ sung  |  ');rr.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text='',bold=False,style=None):
 x=doc.add_paragraph(style=style);r=x.add_run(text);r.bold=bold;return x
def h(text):doc.add_heading(text,level=1)
def h2(text):doc.add_heading(text,level=2)
def newpage(title):doc.add_page_break();h(title)
def f(v):return f'{v:.4f}'
def delta(v):return '≈0' if abs(v)<0.00005 else f'{v:+.4f}'
def val(key,model,metric='AP_mean'):return S[key,model][metric]
def table(headers,rows,widths,font=9.5):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,txt,w in zip(t.rows[0].cells,headers,widths):c.text=str(txt);c.width=Inches(w)
 for row in rows:
  cells=t.add_row().cells
  for c,txt,w in zip(cells,row,widths):c.text=str(txt);c.width=Inches(w)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 t._tbl.tblPr.append(borders)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
  if i==0:repeat=OxmlElement('w:tblHeader');pr.append(repeat)
  for j,c in enumerate(row.cells):
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=c._tc.get_or_add_tcPr()
   margins=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    a=OxmlElement('w:'+side);a.set(qn('w:w'),'75');a.set(qn('w:type'),'dxa');margins.append(a)
   tcPr.append(margins)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6EF' if i==0 else ('F4F7FA' if i%2==0 else 'FFFFFF'));tcPr.append(shade)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1);para.paragraph_format.line_spacing=1.0
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.size=Pt(font);run.bold=(i==0);run.font.color.rgb=RGBColor(0,0,0)
 p('').paragraph_format.space_after=Pt(0)
 return t
def metric_rows(keys):
 return [[label,f(val(k,'LR')),f(val(k,'XGB')),f(val(k,'LR','F1_mean')),f(val(k,'XGB','F1_mean'))] for k,label in keys]

doc.add_paragraph('Kết quả thử nghiệm feature engineering',style='Title')
p('Dự án ADY201m nhóm 6   |   Ngày 29 tháng 10 năm 2026',style='Subtitle')
p('Tài liệu đọc nhanh trước buổi báo cáo sáng 30 tháng 10',bold=True)
p('Báo cáo thử 28 phương án feature trên Logistic Regression và XGBoost bằng 5 fold của Train, tổng cộng 280 lượt huấn luyện. Kết quả giúp giải thích đánh đổi của bộ 13 và 10, đồng thời tìm hướng khai thác bộ 21. Notebook và dữ liệu của dự án được giữ nguyên.')
h2('Kết luận cần nhớ')
p('1. Bộ 21 vẫn là mốc đối chiếu tốt. Trong lượt chạy này, bộ 13 và 10 đều giảm AP ở cả 5 fold, trên cả hai mô hình. Đây là bằng chứng về đánh đổi khi rút gọn trong thiết kế đã thử, chưa chứng minh hai bộ này luôn kém ở mọi cấu hình.')
p('2. Đổi cách biểu diễn có thể hữu ích hơn nén biến. Với Logistic Regression, spline BMI tăng AP từ 0,4015 lên 0,4164 và tăng ở cả 5 fold. F1 tại ngưỡng 0,5 chỉ tăng từ 0,4092 lên 0,4101, nên chưa thể nói mọi chỉ số đều cải thiện rõ.')
p('3. Không nên tuyên bố mọi biến mới đều tốt hoặc mọi biến bị bỏ đều vô ích. Tác động khác nhau theo mô hình; riêng việc bỏ CholCheck làm giảm AP ở cả 5 fold của cả hai mô hình.')
table(['Bộ feature','AP LR','AP XGB','F1 LR','F1 XGB'],metric_rows([('full21','21 gốc'),('compact13','13 của nhóm'),('replacement10','10 của nhóm'),('spline_bmi','21 với spline BMI')]),[2.4,1.05,1.05,1.05,1.05],10)
p('AP và F1 có thang 0 đến 1; càng cao càng tốt. Các số là trung bình của 5 fold. AP tóm tắt chất lượng Precision–Recall qua nhiều ngưỡng; F1 ở đây dùng cố định ngưỡng 0,5.',style='Caption')
p('Đề nghị cho buổi báo cáo: giữ nguyên nội dung nhóm đã học. Dùng kết quả này làm phần thảo luận bổ sung, nói rõ đang kiểm chứng các giả thuyết tạo và rút gọn feature; chưa công bố bộ tối ưu mới.',bold=True)

newpage('Thiết kế thí nghiệm và giới hạn')
p('Dữ liệu BRFSS2015, người 30–59 tuổi; nhãn 1 gộp tiền tiểu đường và tiểu đường tự khai. Chỉ 93.329 dòng Train từ CSV sạch giữ duplicate, gồm 10.192 nhãn 1, tham gia fit và chấm điểm. Dùng split_manifest.csv hiện có; Test không được chấm hoặc dùng chọn phương án.')
table(['Fold','Dòng fit','Dòng validation','Nhãn 1 validation'],[[r['fold'],f"{r['fit_rows']:,}",f"{r['validation_rows']:,}",f"{r['validation_positive']:,}"] for r in D['folds']],[.65,1.7,2,2.25],10)
h2('Các yếu tố được giữ cố định')
p('Dùng đúng 5 fold có sẵn, giữ các profile 21 feature giống nhau cùng fold. Chỉ dùng bản Train giữ duplicate, không NearMiss hoặc SMOTENC. Trọng số lớp tính lại từ phần fit của từng fold. Threshold cố định là xác suất lớn hơn 0,5; không tìm ngưỡng trên validation.')
p('Logistic Regression dùng C=1, lbfgs, max_iter=1500 và class_weight=balanced. XGBoost dùng 100 cây, depth=6, learning_rate=0,1, tree_method=hist và scale_pos_weight=n0/n1 của fold fit. Seed=42. Không tune tham số theo từng bộ feature.')
p('LR chuẩn hóa biến số và one-hot biến thứ bậc. XGB giữ mã thứ bậc gốc. Các trạng thái danh nghĩa mới được one-hot ở cả hai mô hình. Scaler và vị trí nút spline chỉ học từ phần fit, sau đó transform validation.')
h2('Cách đọc bằng chứng')
p('AP là chỉ số chính để khảo sát chất lượng xếp hạng lớp 1; F1, Precision và Recall tại 0,5 bổ sung ý nghĩa phân loại. Delta luôn so với bộ 21 ở cùng mô hình và cùng fold, trừ bảng khôi phục bộ 10 ghi rõ mốc riêng. “≈0” là delta có độ lớn dưới 0,00005, không nhất thiết bằng 0 chính xác.')
p('Kiểm tra độ nhạy bổ sung chỉ chấm một hồ sơ cho mỗi profile 21 feature cộng nhãn trong validation; phần fit vẫn giữ duplicate. Mục đích là xem kết luận có phụ thuộc mạnh vào số lần lặp hồ sơ hay không.')
h2('Giới hạn cần nói rõ')
p('Các ý tưởng đã được hình thành từ EDA trên toàn Train. Vì vậy CV này là kiểm tra phát triển, chưa phải xác nhận độc lập của toàn bộ quá trình lựa chọn. 28 phương án dùng chung các fold; việc xem nhiều kết quả có thể làm cực đại lạc quan. Độ lệch chuẩn giữa fold không phải khoảng tin cậy và số fold thắng không phải kiểm định ý nghĩa thống kê.')
p('Chưa thử Random Forest, Decision Tree, tuning, dữ liệu ngoài hoặc trọng số khảo sát. Các điểm ở đây không so trực tiếp với Test trong notebook vì tập đánh giá khác. Các nhóm vẫn được xác định theo 21 feature gốc, không chia lại theo bộ rút gọn.')

newpage('Bộ 21 còn có thể khai thác tốt hơn')
p('Hai hướng nổi bật của LR là biểu diễn phi tuyến BMI bằng spline và thêm tổ hợp tuổi × nhóm BMI. Đây đều là thay đổi cách biểu diễn trên nền thông tin gốc, không yêu cầu bỏ Age hoặc BMI để chỉ giữ một tích.')
doc.add_picture(str(P/'overview_ap.png'),width=Inches(6.65))
p('Hình 1  AP trung bình và độ lệch chuẩn giữa 5 fold. Đường đứt là AP của 21 gốc ở đúng mô hình. Thanh sai số thể hiện độ dao động giữa fold, không phải khoảng tin cậy. Trục AP chỉ hiển thị khoảng 0,375 đến 0,440.',style='Caption')
table(['Phương án','ΔAP LR','ΔAP XGB','ΔF1 LR','ΔF1 XGB'],[[label,delta(val(k,'LR','Delta_AP_mean')),delta(val(k,'XGB','Delta_AP_mean')),delta(val(k,'LR','Delta_F1_mean')),delta(val(k,'XGB','Delta_F1_mean'))] for k,label in [('spline_bmi','Spline BMI'),('age_bmi_grid','Tuổi × nhóm BMI'),('education_income_state','Học vấn × thu nhập'),('day_flags','Thêm hai cờ số ngày'),('log_days','Log1p hai số ngày')]],[2.7,.975,.975,.975,.975],9.5)
p('Spline BMI dùng 5 nút phân vị, bậc 3 và 6 cột cơ sở thay một cột BMI. LR có thể học một đường cong theo BMI thay vì một hệ số tuyến tính. AP tăng lần lượt +0,0124; +0,0118; +0,0172; +0,0184; +0,0148 ở 5 fold. F1 chỉ tăng ở 3/5 fold.')
p('Tổ hợp học vấn × thu nhập cho XGB có ΔAP=+0,0034, thắng 5/5 fold; LR lại giảm AP. Vì thế chưa có một cách biểu diễn tốt chung cho mọi mô hình. Cờ số ngày và log1p không đổi kết quả XGB trong lần chạy này, phù hợp với việc mô hình cây đã có thể tách các ngưỡng trên biến gốc.')
p('Ở validation bỏ trùng, spline BMI của LR vẫn tăng AP khoảng +0,0150 so với bộ 21. Đây là độ nhạy bổ sung, không phải một lần xác nhận trên nguồn dữ liệu mới.')

newpage('Kiểm chứng riêng bốn biến mới của nhóm')
p('Mỗi dòng dưới chỉ thay đổi một ý tưởng so với bộ 21. “Thêm” giữ các biến gốc; “thay” bỏ các biến thành phần được chỉ định rồi đưa biến mới vào. Các delta là trung bình 5 chênh lệch theo fold.')
rows=[]
for name,label in [('AgeBMI','AgeBMI'),('BPCholState','BPCholState'),('LowEducationLowIncome','LowEducationLowIncome'),('PoorHealthDiffWalk','PoorHealthDiffWalk')]:
 for prefix,word in [('add_','Thêm '),('replace_','Thay bằng ')]:
  k=prefix+name;rows.append([word+label,delta(val(k,'LR','Delta_AP_mean')),delta(val(k,'XGB','Delta_AP_mean')),delta(val(k,'LR','Delta_F1_mean')),delta(val(k,'XGB','Delta_F1_mean'))])
table(['Thay đổi so với 21','ΔAP LR','ΔAP XGB','ΔF1 LR','ΔF1 XGB'],rows,[2.7,.975,.975,.975,.975],9.5)
h2('AgeBMI')
p('Định nghĩa (Age×5+17)×BMI, với Age là mã nhóm 3–8. Khi thêm vào bộ 21, AP tăng nhẹ ở cả hai mô hình. Khi thay Age và BMI, AP của LR gần mốc 21 nhưng XGB giảm khoảng 0,0010. Việc nén thực sự mất thông tin, nhưng mức thay đổi metrics còn phụ thuộc mô hình và regularization; chưa thể khẳng định thay tích luôn làm điểm giảm.')
h2('BPCholState')
p('Định nghĩa 2×HighBP+HighChol, gồm bốn trạng thái danh nghĩa. Thay hai biến gốc bằng trạng thái giữ đủ thông tin của cặp nhị phân. LR gần như không đổi AP; XGB tăng nhẹ trong lần thử. Đây là đổi biểu diễn, không phải có thêm dữ liệu độc lập.')
h2('LowEducationLowIncome')
p('Định nghĩa Education≤3 và Income≤3. Thêm cờ gần như không giúp LR; thay cả Education/Income làm F1 giảm khoảng 0,0042 ở LR và 0,0031 ở XGB. Một cờ nhị phân chưa biểu diễn đủ sự khác biệt giữa các mức học vấn và thu nhập ở thiết kế này.')
h2('PoorHealthDiffWalk')
p('Định nghĩa GenHlth≥4 và DiffWalk=1. Phương án “thay” chỉ bỏ DiffWalk, vẫn giữ GenHlth. Tác động nhỏ và khác nhau theo mô hình. Chưa đủ bằng chứng để ưu tiên phép thay này hoặc gọi biến kết hợp là quan trọng hơn biến nguồn.')
p('Bộ 25 thêm cả bốn biến có ΔAP LR=+0,0008 và XGB=+0,0017. Không thể quy toàn bộ phần tăng cho một biến riêng, và mức tăng này chưa xác nhận lợi ích độc lập.',bold=True)

newpage('Bộ 10 mất thông tin ở đâu')
p('Để kiểm tra trực tiếp bộ 10, ba trường hợp dưới khôi phục một cụm biến nguồn và bỏ phép nén tương ứng. Các feature còn lại của bộ 10 giữ nguyên. Bảng này dùng mốc bộ 10, không dùng mốc bộ 21.')
restore=[('restore10_AgeBMI','Khôi phục Age và BMI'),('restore10_LowEducationLowIncome','Khôi phục Education và Income'),('restore10_PoorHealthDiffWalk','Khôi phục DiffWalk')]
table(['Thay đổi từ bộ 10','ΔAP LR','ΔAP XGB','ΔF1 LR','ΔF1 XGB'],[[label,delta(val(k,'LR')-val('replacement10','LR')),delta(val(k,'XGB')-val('replacement10','XGB')),delta(val(k,'LR','F1_mean')-val('replacement10','LR','F1_mean')),delta(val(k,'XGB','F1_mean')-val('replacement10','XGB','F1_mean'))] for k,label in restore],[2.7,.975,.975,.975,.975],9.5)
p('Khôi phục Education và Income tăng AP và F1 ở cả hai mô hình. Đây là dấu hiệu trực tiếp nhất trong ba phép khôi phục rằng cờ LowEducationLowIncome đã nén quá nhiều thông tin hữu ích đối với bộ 10 hiện tại.')
p('Khôi phục Age/BMI chỉ tăng nhẹ ở XGB và giảm ở LR. Kết quả này nhắc rằng giữ thêm thông tin không tự động làm mô hình cố định tốt hơn. Muốn kết luận đầy đủ cần xem biểu diễn, regularization và validation phù hợp.')
h2('Số feature không bằng số câu hỏi cần thu thập')
p('Bộ 10 cần 13 biến gốc: Stroke, HeartDiseaseorAttack, PhysActivity, GenHlth, MentHlth, PhysHlth, Age, BMI, HighBP, HighChol, Education, Income và DiffWalk. Vì vậy không nên gọi đây là bộ chỉ cần 10 câu trả lời gốc.')
p('Sau mã hóa, bộ 10 tạo 15 cột cho LR và 12 cột cho XGB. Bộ 21 tạo 38 cột cho LR và 21 cột cho XGB. Cần phân biệt biến nguồn, feature trước mã hóa và cột mô hình sau mã hóa khi trình bày hiệu quả rút gọn.')
h2('Vì sao không tự sửa bộ 10 ngay')
p('Lượt này chỉ khôi phục từng cụm, chưa kiểm chứng một bộ kết hợp tối ưu mới. Nếu ghép ngay tất cả lựa chọn có điểm tốt, nhóm sẽ có một thiết kế chưa được đánh giá như một chỉnh thể. Báo cáo ngày mai nên trình bày các giới hạn của bộ 10 và hướng kiểm chứng tiếp, giữ nguyên phương án nhóm đã chuẩn bị.')
h2('Định nghĩa các biểu diễn bổ sung')
p('Cờ số ngày là MentHlth>0 và PhysHlth>0, giữ hai số ngày gốc. Log1p thay hai số ngày bằng ln(1+x). Tổ hợp tuổi × BMI gồm 6 nhóm tuổi và 4 nhóm BMI <18,5; 18,5–<25; 25–<30; ≥30. Tổ hợp học vấn × thu nhập gồm 6×8 trạng thái. Hai tổ hợp đều được thêm vào 21 gốc rồi one-hot; không giả định mã trạng thái là số có khoảng cách tuyến tính.')

newpage('Bộ 13 thay đổi thế nào so với bộ 21')
p('Bộ 13 giữ nguyên 13 biến gốc, bỏ 8 biến khỏi bộ 21 và không tạo biến kết hợp. Bảng dưới lấy bộ 21 làm mốc: delta bằng điểm bộ 13 trừ điểm bộ 21, trên cùng mô hình và cùng 5 fold. Số âm là giảm; Precision, Recall và F1 dùng ngưỡng 0,5.')
table(['Bộ 13 so với 21','ΔAP','ΔF1','ΔPrecision','ΔRecall'],[[model,delta(val('compact13',model)-val('full21',model)),delta(val('compact13',model,'F1_mean')-val('full21',model,'F1_mean')),delta(val('compact13',model,'Precision_mean')-val('full21',model,'Precision_mean')),delta(val('compact13',model,'Recall_mean')-val('full21',model,'Recall_mean'))] for model in ['LR','XGB']],[2.4,1.05,1.05,1.05,1.05],10)
p('LR: AP giảm 0,4015 xuống 0,3971; F1 giảm 0,4092 xuống 0,4061. XGB: AP giảm 0,4076 xuống 0,4027; F1 giảm 0,4096 xuống 0,4053. AP giảm ở cả 5 fold của mỗi mô hình; F1 giảm ở 4/5 fold LR và 5/5 fold XGB. Đây là giảm điểm trên thang 0–1, không phải phần trăm giảm tương đối.')
h2('Tám biến bị bỏ có hoàn toàn vô ích không')
p('Các phép bỏ riêng từng biến khỏi bộ 21 giúp kiểm tra căn cứ rút gọn. Bảng này vẫn lấy bộ 21 làm mốc; đây không phải thí nghiệm thêm lại từng biến vào bộ 13.')
drop_keys=[x['key'] for x in protocol['experiments'] if x['kind']=='drop']
table(['Bỏ riêng khỏi 21','ΔAP LR','ΔAP XGB','ΔF1 LR','ΔF1 XGB'],[[k[5:],delta(val(k,'LR','Delta_AP_mean')),delta(val(k,'XGB','Delta_AP_mean')),delta(val(k,'LR','Delta_F1_mean')),delta(val(k,'XGB','Delta_F1_mean'))] for k in drop_keys],[2.7,.975,.975,.975,.975],9.5)
p('Bỏ CholCheck làm AP giảm ở cả 5 fold của cả hai mô hình; bỏ HvyAlcoholConsump cũng giảm AP trung bình ở cả hai. Smoker có tác động trái dấu theo mô hình. Vì vậy, liên hệ đơn biến nhỏ chưa đủ chứng minh một feature không có giá trị bổ sung. Không cộng các delta riêng lẻ để suy ra tác động của việc bỏ cả tám biến.')
h2('Đánh đổi của bộ 13')
p('Bộ 13 giữ HighBP, HighChol, BMI, Stroke, HeartDiseaseorAttack, PhysActivity, GenHlth, MentHlth, PhysHlth, DiffWalk, Age, Education và Income. Số biến nguồn giảm 21 xuống 13; sau mã hóa LR giảm 38 xuống 30 cột, XGB giảm 21 xuống 13 cột. Đổi lại, AP và F1 trung bình giảm trong cấu hình đã thử.')
p('Cách trình bày: “Bộ 13 là phương án rút gọn có đánh đổi, chưa vượt bộ 21. Nhóm cần kiểm chứng từng quyết định loại biến thay vì chỉ dùng tương quan hoặc p-value. Kết quả này là kiểm tra phát triển trên Train, chưa xác nhận một bộ tối ưu mới.”',bold=True)

newpage('Bỏ từng biến và cách trình bày ngày mai')
p('Tám biến từng bị bỏ khi chuyển từ 21 xuống 13 được kiểm tra riêng. Giá trị âm cho biết bỏ biến làm AP giảm; dương cho biết AP tăng trong cấu hình và fold đã thử.')
doc.add_picture(str(P/'drop_one_ap.png'),width=Inches(6.65))
p('Hình 2  Chênh lệch AP khi bỏ riêng từng biến khỏi bộ 21. Hai màu tương ứng hai mô hình. Đây là thí nghiệm bỏ biến và huấn luyện lại, không phải xếp hạng tương quan hoặc bằng chứng nhân quả.',style='Caption')
p('CholCheck đáng chú ý nhất: bỏ riêng biến này giảm AP khoảng 0,0025 ở LR và 0,0026 ở XGB, giảm ở cả 5 fold của hai mô hình. HvyAlcoholConsump cũng giảm AP trung bình ở cả hai mô hình. Ngược lại, Smoker làm AP thay đổi trái dấu giữa LR và XGB. Không nên coi các biến có liên hệ đơn biến nhỏ là chắc chắn vô ích.')
p('Không cộng các delta bỏ từng biến để dự đoán điểm khi bỏ đồng thời tám biến, vì chúng có thể bổ sung hoặc thay thế thông tin cho nhau. Bộ 13 đã được chạy trực tiếp và giảm AP khoảng 0,0044 ở LR, 0,0049 ở XGB so với 21.')
h2('Đoạn trình bày có thể dùng trong khoảng một phút')
p('“Nhóm xây dựng các bộ rút gọn như những giả thuyết cần kiểm chứng. Một thử nghiệm bổ sung giữ nguyên dự án và dùng 5 fold của Train cho thấy bộ 21 vẫn có giá trị làm mốc. Bộ 13 và 10 có đánh đổi về thông tin; riêng cờ học vấn thấp và thu nhập thấp chưa thay thế tốt hai biến nguồn. Hướng phát triển đáng chú ý là đổi cách biểu diễn, chẳng hạn spline BMI giúp Logistic Regression tăng AP ổn định qua 5 fold, trong khi F1 ở ngưỡng 0,5 chỉ tăng ít. Do các ý tưởng được hình thành từ EDA trên Train, nhóm xem đây là kết quả phát triển, chưa công bố bộ tối ưu được xác nhận độc lập.”')
h2('Nếu giảng viên hỏi vì sao không dùng bộ mới ngay')
p('Trả lời: thí nghiệm bổ sung dùng để hiểu đánh đổi của feature. Nhóm chưa khóa và xác nhận một cấu hình mới bằng dữ liệu chưa tham gia lựa chọn; do đó báo cáo chính giữ nguyên và phần này được trình bày như hướng cải tiến.')
p('Nếu hỏi “spline có dự báo tốt hơn không”, trả lời cụ thể: AP của LR tăng 0,0149 trung bình nhưng F1 tại 0,5 chỉ tăng 0,0009. Không khẳng định tốt hơn ở mọi ngưỡng, mọi mô hình hoặc nguồn dữ liệu khác.')

labels={
 'full21':'21 biến gốc','compact13':'13 biến của nhóm','replacement10':'10 biến của nhóm','expanded25':'21 cộng cả 4 biến mới',
 'add_AgeBMI':'Thêm AgeBMI','replace_AgeBMI':'Thay Age và BMI bằng AgeBMI',
 'add_BPCholState':'Thêm BPCholState','replace_BPCholState':'Thay HighBP và HighChol',
 'add_LowEducationLowIncome':'Thêm cờ học vấn và thu nhập thấp','replace_LowEducationLowIncome':'Thay Education và Income bằng cờ',
 'add_PoorHealthDiffWalk':'Thêm cờ sức khỏe kém và khó đi','replace_PoorHealthDiffWalk':'Thay DiffWalk bằng cờ kết hợp',
 'restore10_AgeBMI':'Bộ 10 khôi phục Age và BMI','restore10_LowEducationLowIncome':'Bộ 10 khôi phục Education và Income','restore10_PoorHealthDiffWalk':'Bộ 10 khôi phục DiffWalk',
 'day_flags':'Thêm hai cờ có ngày sức khỏe kém','log_days':'Thay hai số ngày bằng log1p','spline_bmi':'Thay BMI bằng spline bậc 3',
 'age_bmi_grid':'Thêm tổ hợp tuổi và nhóm BMI','education_income_state':'Thêm tổ hợp học vấn và thu nhập'}
for spec in protocol['experiments']:
 if spec['key'].startswith('drop_'):labels[spec['key']]='Bỏ riêng '+spec['key'][5:]

for page,start in enumerate([0,14]):
 newpage('Bảng kết quả đầy đủ '+('phần một' if page==0 else 'phần hai'))
 p('AP và F1 là trung bình 5 fold trên validation giữ duplicate. “Cột” là số feature trước mã hóa; riêng spline BMI bắt đầu từ 21 biến nguồn rồi đổi BMI thành nhiều cột cơ sở. Mỗi dòng đã chạy 10 fit gồm 2 mô hình × 5 fold.')
 rr=[]
 for i,spec in enumerate(protocol['experiments'][start:start+14],start+1):
  k=spec['key'];rr.append([str(i),labels[k],str(len(spec['columns'])),f(val(k,'LR')),f(val(k,'XGB')),f(val(k,'LR','F1_mean')),f(val(k,'XGB','F1_mean'))])
 table(['STT','Phương án','Cột','AP LR','AP XGB','F1 LR','F1 XGB'],rr,[.35,2.8,.35,.7,.7,.7,.7],9.5)
 if page==0:
  p('Bảng tiếp theo gồm trường hợp khôi phục DiffWalk, các biểu diễn bổ sung và tám phép bỏ riêng từng biến.',style='Caption')
  h2('Cách diễn giải số nhỏ')
  p('Delta khoảng 0,0001 chỉ là một phần mười nghìn trên thang điểm 0–1. Không nên dùng thứ hạng chênh lệch rất nhỏ để khẳng định ưu thế chắc chắn. Phải xét dấu chênh lệch ở các fold, mục tiêu chỉ số và chi phí thu thập thông tin.')
 else:
  h2('Dấu vết kiểm chứng')
  v=protocol['software']
  p('Đã đối chiếu lại 56 file dự đoán ngoài fold với đủ 280 dòng metrics; không có cảnh báo hội tụ được ghi nhận. SHA 256 của 10 file Git theo dõi trước và sau chạy giống nhau. Không chấm Test. Mã chạy, protocol, metrics từng fold và dự đoán được lưu riêng ngoài repository.',style=None)
  p('Commit nguồn '+protocol['project_commit'][:12]+'. Python '+v['python']+'; scikit-learn '+v['sklearn']+'; XGBoost '+v['xgboost']+'.',style='Caption')
  p('Nguồn phương pháp: scikit-learn Common pitfalls and recommended practices; SplineTransformer; Average precision score. Tài liệu chính thức: https://scikit-learn.org/stable/common_pitfalls.html ; https://scikit-learn.org/stable/modules/generated/sklearn.preprocessing.SplineTransformer.html ; https://scikit-learn.org/stable/modules/generated/sklearn.metrics.average_precision_score.html',style='Caption')

doc.core_properties.title='Kết quả thử nghiệm feature engineering dự án ADY201m'
doc.core_properties.subject='28 phương án feature trên 5 fold Train của BRFSS2015 tuổi 30 đến 59'
doc.core_properties.author='Báo cáo thí nghiệm bổ sung'
out=P/'Bao_cao_thu_nghiem_feature_ADY201m.docx'
doc.save(out)
print(out)
print('Paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'images',len(doc.inline_shapes))
